"""공식 유형명만 읽는다. 코드를 읽거나 생성하지 않는다."""
import re

from openpyxl import load_workbook


def normalize_type(value):
    # 서로 다른 가운뎃점과 공백만 통일한다. 유사어/부분 일치는 허용하지 않는다.
    return re.sub(r"\s+", "", value).translate(str.maketrans({"ㆍ": "·", "‧": "·"}))


def load_reference(path, raw_type):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "교육유형" not in workbook.sheetnames:
            raise ValueError("기준 양식에 교육유형 시트가 없습니다.")
        defined = workbook.defined_names.get("교육유형명")
        if defined is None:
            raise ValueError("기준 양식에 교육유형명 이름정의가 없습니다.")
        official_names = []
        for sheet, cell_range in defined.destinations:
            if sheet != "교육유형":
                raise ValueError("교육유형명 이름정의가 교육유형 시트를 참조하지 않습니다.")
            for row in workbook[sheet][cell_range]:
                official_names.extend(c.value for c in row if isinstance(c.value, str) and c.value.strip())
        matches = [name for name in official_names if normalize_type(name) == normalize_type(raw_type)]
        return {"official": matches[0] if len(matches) == 1 else None,
                "matches": matches, "count": len(official_names)}
    finally:
        workbook.close()
