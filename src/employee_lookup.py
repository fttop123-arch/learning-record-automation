"""직원 Excel의 부서 유무에 따라 부서+성명 또는 이름으로 정확히 조회한다."""
from pathlib import Path

from openpyxl import load_workbook

EMPLOYEE_COLUMNS = {
    "department": "부서",
    "name": "성명",
    "employee_id": "개인식별번호",
}
NAME_COLUMN_ALIASES = ("성명", "이름")


def key_text(value):
    # 양끝 공백만 제거한다. 내부 공백, 유사어, 부서 별칭은 추측하지 않는다.
    return value.strip() if isinstance(value, str) else ""


class EmployeeLookup:
    def __init__(self, path=None, columns=None):
        self.records = {}
        self.match_mode = "no_file"
        self.no_file = path is None or not Path(path).exists()
        if self.no_file:
            return
        default_columns = columns is None
        columns = EMPLOYEE_COLUMNS if default_columns else columns
        if set(columns) != set(EMPLOYEE_COLUMNS) or len(set(columns.values())) != 3:
            raise ValueError("직원정보 컬럼 매핑에는 서로 다른 부서/성명/개인식별번호가 필요합니다.")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            headers = []
            for sheet in workbook:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    labels = [key_text(value) for value in row]
                    name_labels = NAME_COLUMN_ALIASES if default_columns else (columns["name"],)
                    name_indexes = [i for i, label in enumerate(labels) if label in name_labels]
                    if not name_indexes or columns["employee_id"] not in labels:
                        continue
                    if len(name_indexes) != 1 or labels.count(columns["employee_id"]) != 1 or labels.count(columns["department"]) > 1:
                        raise ValueError("직원정보 헤더가 중복되어 컬럼을 확정할 수 없습니다.")
                    indexes = {"name": name_indexes[0], "employee_id": labels.index(columns["employee_id"])}
                    if columns["department"] in labels:
                        indexes["department"] = labels.index(columns["department"])
                    headers.append((sheet.title, row_number, indexes))
            if len(headers) != 1:
                raise ValueError("직원정보의 헤더를 하나로 확인할 수 없습니다. 컬럼 매핑과 시트를 확인해주세요.")
            sheet_name, header_row, indexes = headers[0]
            self.match_mode = "department_name" if "department" in indexes else "name_only"
            for row_number, row in enumerate(workbook[sheet_name].iter_rows(min_row=header_row + 1), header_row + 1):
                department = key_text(row[indexes["department"]].value) if "department" in indexes else None
                name = key_text(row[indexes["name"]].value)
                if not name or (self.match_mode == "department_name" and not department):
                    continue
                cell = row[indexes["employee_id"]]
                employee_id = key_text(cell.value) if cell.data_type not in ("f", "e") else ""
                record = {"employee_id": employee_id or None, "sheet": sheet_name, "row": row_number}
                key = (department, name) if self.match_mode == "department_name" else name
                self.records.setdefault(key, []).append(record)
        finally:
            workbook.close()

    def lookup(self, department, name):
        department, name = key_text(department), key_text(name)
        result = {"department": department, "name": name, "employee_id": None, "match_mode": self.match_mode}
        if self.no_file:
            return dict(result, status="no_file", reason="내부 직원정보 파일 미지정 또는 없음")
        key = (department, name) if self.match_mode == "department_name" else name
        candidates = self.records.get(key, []) if name and (department or self.match_mode == "name_only") else []
        if len(candidates) > 1:
            reason = "동명이인 확인 필요: " + ("같은 부서에 같은 성명 2명 이상" if self.match_mode == "department_name" else "직원정보 파일에 같은 이름 2명 이상")
            return dict(result, status="duplicate", reason=reason,
                        candidates=[{"sheet": c["sheet"], "row": c["row"]} for c in candidates])
        if not candidates:
            return dict(result, status="not_found", reason="내부 직원정보의 대상 부서에서 찾지 못함" if self.match_mode == "department_name" else "직원정보 파일에서 이름을 찾지 못함")
        if candidates[0]["employee_id"] is None:
            return dict(result, status="missing_id", reason="개인식별번호 빈셀 또는 확인 필요: 텍스트 상수 값 필요")
        return dict(result, status="matched", employee_id=candidates[0]["employee_id"], reason="부서+성명 정확히 일치" if self.match_mode == "department_name" else "부서 선별 후 이름 정확히 일치 (직원정보에 부서 컬럼 없음)")


def matched_employee_id(attendee):
    """해당 대상자에 묶인 matched 결과만 B열에 전달한다."""
    result = attendee.get("employee_lookup", {})
    if result.get("status") != "matched":
        return None
    if (result.get("department"), result.get("name")) != (key_text(attendee["department"]), key_text(attendee["name"])):
        raise ValueError("직원정보 조회 결과와 등록 대상자가 일치하지 않습니다.")
    value = result.get("employee_id")
    if not isinstance(value, str) or not value.strip():
        raise ValueError("matched 결과에 유효한 개인식별번호가 없습니다.")
    return value
