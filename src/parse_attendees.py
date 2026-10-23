"""헤더를 찾아 대상 부서의 참석자와 부서 내 성명 중복을 확인한다."""
from collections import Counter
import re

from openpyxl import load_workbook

TARGET_DEPARTMENT = "디지털정책과"


def normalize_space(value):
    return re.sub(r"\s+", "", str(value or ""))


def parse_attendees(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        candidates = []
        for sheet in workbook:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                labels = [normalize_space(v) for v in row]
                if all(labels.count(label) == 1 for label in ("소속", "직급", "성명")):
                    candidates.append((sheet.title, row_number, {label: labels.index(label) for label in ("소속", "직급", "성명")}))
        if len(candidates) != 1:
            raise ValueError("소속/직급/성명 헤더를 가진 명단을 하나로 확인할 수 없습니다.")
        sheet_name, header_row, columns = candidates[0]
        attendees, total = [], 0
        for row_number, row in enumerate(workbook[sheet_name].iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            values = {label: str(row[index] or "").strip() for label, index in columns.items()}
            if not any(values.values()):
                continue
            total += 1
            if normalize_space(values["소속"]) != TARGET_DEPARTMENT:
                continue
            attendees.append({"department": TARGET_DEPARTMENT, "rank": values["직급"], "name": values["성명"], "row": row_number})
        counts = Counter(normalize_space(a["name"]) for a in attendees if a["name"])
        for attendee in attendees:
            attendee["status"] = (
                "성명 누락: 사용자 확인 필요" if not attendee["name"] else
                "동명이인 확인 필요" if counts[normalize_space(attendee["name"])] > 1 else
                "직급 누락: 사용자 확인 필요" if not attendee["rank"] else "대상 부서 일치"
            )
        return {"sheet": sheet_name, "total": total, "attendees": attendees}
    finally:
        workbook.close()
