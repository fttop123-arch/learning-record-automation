"""XLSX 패키지를 복제하고 Sheet1 데이터 및 계산 의존 정보만 변경한다."""
from datetime import date
from io import BytesIO
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import to_excel

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHEET = "xl/worksheets/sheet1.xml"


def ask_choice(label, choices, default=None):
    while True:
        suffix = f" [기본값 {default}]" if default is not None else " [반드시 선택]"
        value = input(f"{label} ({'/'.join(choices)}){suffix}: ").strip().upper()
        if not value and default is not None:
            value = default
        if value in choices:
            return value
        print(f"잘못된 입력입니다. {' 또는 '.join(choices)} 중 선택해주세요.")


def validate_ready(roster, reference, method_code):
    if not roster["attendees"]:
        raise ValueError("대상자가 없습니다.")
    if any(a["status"] != "대상 부서 일치" for a in roster["attendees"]):
        raise ValueError("동명이인 또는 누락 항목 확인 후 생성할 수 있습니다.")
    if not reference["official"] or method_code is None:
        raise ValueError("교육유형명 또는 학습방법이 미확정입니다.")


def build_excel(template, notice, attendees, official, required, welfare, participation):
    if required not in ("N", "Y") or participation not in ("0", "1"):
        raise ValueError("사용자 선택값이 올바르지 않습니다.")
    if welfare != "N":
        raise ValueError("사회복지교육=Y: 시간/분 추가 확인 필요. 결과 파일을 생성하지 않습니다.")
    with ZipFile(template) as source:
        original = source.read(SHEET).decode("utf-8")
        root = ET.fromstring(original)
        sample = root.find('m:sheetData/m:row[@r="2"]', NS)
        if sample is None:
            raise ValueError("기준 양식 2행이 없습니다.")
        existing_rows = root.findall('m:sheetData/m:row', NS)
        if any(int(row.get("r")) > 2 for row in existing_rows):
            raise ValueError("기준 양식에 예상하지 못한 추가 행이 있어 중단합니다.")
        formula = sample.find('m:c[@r="D2"]/m:f', NS)
        if formula is None or not formula.text:
            raise ValueError("기준 양식 D2 수식이 없습니다.")
        styles = {c.get("r")[:-1]: c.get("s") for c in sample}
        serial = to_excel(date.fromisoformat(notice["date"]))
        generated_rows = []
        for index, attendee in enumerate(attendees, 2):
            values = [attendee["name"], None, official, None, required, serial, serial,
                      notice["hours"], notice["minutes"], 1, int(participation),
                      notice["name"], notice["name"], welfare, None, None]
            row_tag = re.search(r'<row\b[^>]*\br="2"[^>]*>', original).group()
            row_tag = re.sub(r'\br="2"', f'r="{index}"', row_tag, count=1)
            cells = []
            for column, value in zip("ABCDEFGHIJKLMNOP", values):
                ref = f"{column}{index}"
                style = f' s="{styles[column]}"' if styles.get(column) is not None else ""
                if column == "D":
                    translated = Translator("=" + formula.text, origin="D2").translate_formula(ref)[1:]
                    cells.append(f'<c r="{ref}"{style}><f>{escape(translated)}</f></c>')
                elif value is None:
                    cells.append(f'<c r="{ref}"{style}/>')
                elif isinstance(value, (int, float)):
                    cells.append(f'<c r="{ref}"{style}><v>{value}</v></c>')
                else:
                    cells.append(f'<c r="{ref}"{style} t="inlineStr"><is><t>{escape(value)}</t></is></c>')
            generated_rows.append(f'{row_tag}{"".join(cells)}</row>')
        # 기존 네임스페이스, 유효성검사 및 서식 XML은 원문 그대로 보존한다.
        header = re.search(r'<row\b[^>]*\br="1"[^>]*>.*?</row>', original, re.S)
        if header is None:
            raise ValueError("기준 양식의 머리글을 확인할 수 없습니다.")
        updated = re.sub(r'<sheetData\b[^>]*>.*?</sheetData>',
                         lambda _: '<sheetData>' + header.group() + ''.join(generated_rows) + '</sheetData>', original, count=1, flags=re.S)
        updated = re.sub(r'<dimension\b[^>]*/>', f'<dimension ref="A1:P{len(attendees)+1}"/>', updated, count=1)
        buffer = BytesIO()
        with ZipFile(buffer, "w") as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename == SHEET:
                    payload = updated.encode("utf-8")
                elif entry.filename == "xl/workbook.xml":
                    text = payload.decode("utf-8")
                    text = re.sub(r'<calcPr\b[^>]*/>', '<calcPr calcId="191029" fullCalcOnLoad="1" forceFullCalc="1"/>', text, count=1)
                    payload = text.encode("utf-8")
                elif entry.filename == "xl/calcChain.xml":
                    # 복사한 D열 수식도 계산 의존 목록에 포함한다.
                    text = payload.decode("utf-8")
                    text = re.sub(r'<calcChain\b[^>]*>.*?</calcChain>',
                                  lambda m: m.group().split('>', 1)[0] + '>' + ''.join(f'<c r="D{i}" i="1"/>' for i in range(2, len(attendees)+2)) + '</calcChain>', text, flags=re.S)
                    payload = text.encode("utf-8")
                target.writestr(entry, payload)
        return buffer.getvalue()


def verify_excel(template, payload, notice, attendees, official, required, participation):
    with ZipFile(template) as source, ZipFile(BytesIO(payload)) as result:
        if source.namelist() != result.namelist():
            raise ValueError("검증 실패: XLSX 구성요소 변경")
        allowed = {SHEET, "xl/workbook.xml", "xl/calcChain.xml"}
        for name in source.namelist():
            if name not in allowed and source.read(name) != result.read(name):
                raise ValueError(f"검증 실패: 보존 대상 변경 {name}")
        before = ET.fromstring(source.read(SHEET))
        after = ET.fromstring(result.read(SHEET))
        for tree in (before, after):
            for tag in ("sheetData", "dimension"):
                element = tree.find("m:" + tag, NS)
                if element is not None:
                    tree.remove(element)
        if ET.tostring(before) != ET.tostring(after):
            raise ValueError("검증 실패: 데이터 이외의 Sheet1 설정 변경")
    workbook = load_workbook(BytesIO(payload), data_only=False)
    try:
        sheet = workbook["Sheet1"]
        if sheet.max_row != len(attendees) + 1:
            raise ValueError("검증 실패: 실제 데이터 행 수")
        for row, attendee in enumerate(attendees, 2):
            expected = [attendee["name"], None, official,
                        f'=INDIRECT("유형코드"&MATCH(C{row},교육유형명,0))', required,
                        date.fromisoformat(notice["date"]), date.fromisoformat(notice["date"]),
                        notice["hours"], notice["minutes"], 1, int(participation),
                        notice["name"], notice["name"], "N", None, None]
            for column, value in enumerate(expected, 1):
                actual = sheet.cell(row, column).value
                if column in (6, 7) and hasattr(actual, "date"):
                    actual = actual.date()
                if actual != value:
                    raise ValueError(f"검증 실패: {sheet.cell(row, column).coordinate}")
    finally:
        workbook.close()


def confirm_and_generate(template, output_dir, notice, roster, reference, method_code):
    validate_ready(roster, reference, method_code)
    print("\n=== 사용자 확인값 입력 ===")
    required = ask_choice("필수교육 여부", ("N", "Y"), "N")
    welfare = ask_choice("사회복지교육 여부", ("N", "Y"), "N")
    print("참여유형: 0 = 오프라인교육(대면), 1 = 온라인교육(비대면)")
    print("이번 샘플을 대면교육으로 확인하셨다면 0을 선택해주세요.")
    participation = ask_choice("참여유형", ("0", "1"))
    payload = build_excel(template, notice, roster["attendees"], reference["official"], required, welfare, participation)
    verify_excel(template, payload, notice, roster["attendees"], reference["official"], required, participation)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', notice["name"]).replace(' ', '_')
    destination = output_dir / f"{stem}_상시학습등록.xlsx"
    index = 2
    while True:
        try:
            with destination.open("xb") as stream:
                stream.write(payload)
            break
        except FileExistsError:
            destination = output_dir / f"{stem}_상시학습등록_{index}.xlsx"
            index += 1
    # 저장된 실제 파일도 다시 검증한다.
    verify_excel(template, destination.read_bytes(), notice, roster["attendees"], reference["official"], required, participation)
    print(f"\n결과 Excel 생성: {destination.resolve()}")
    print(f"검증 완료: 대상자 {len(roster['attendees'])}명, 개인식별번호 빈셀, 행별 수식, 입력값 및 양식 보존")
    print("개인식별번호가 비어 있으므로 내부자료 연계 전 등록 완료본이 아닙니다.")
    print("D열 수식 결과는 Excel에서 열 때 재계산됩니다. 현재 검증은 수식과 참조 구조 검증입니다.")
