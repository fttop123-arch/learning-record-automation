"""가상 Excel만으로 조회 및 기존 생성 경로를 검증한다."""
from contextlib import redirect_stdout
from copy import deepcopy
from io import BytesIO, StringIO
from pathlib import Path
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from employee_lookup import EmployeeLookup, matched_employee_id
from generate_excel import build_excel, verify_excel, confirm_and_generate
from parse_notice import parse_notice
from parse_attendees import parse_attendees
from load_reference import load_reference

FIXTURES = ROOT / "tests" / "fixtures"


def two_column_lookup(rows):
    """파일 저장 없이 두 컬럼 시트의 읽기 결과를 주입한다."""
    sheet = MagicMock()
    sheet.title = "직원정보"
    def iter_rows(min_row=1, values_only=False):
        selected = rows[min_row - 1:]
        if values_only:
            return iter(selected)
        return iter([tuple(SimpleNamespace(value=v, data_type="s" if v is not None else "n") for v in row) for row in selected])
    sheet.iter_rows.side_effect = iter_rows
    workbook = MagicMock()
    workbook.__iter__.side_effect = lambda: iter([sheet])
    workbook.__getitem__.return_value = sheet
    with patch("employee_lookup.load_workbook", return_value=workbook):
        return EmployeeLookup(FIXTURES / "employees.xlsx")


class TwoColumnEmployeeTests(unittest.TestCase):
    def test_name_aliases_and_exact_match(self):
        for label in ("이름", "성명"):
            lookup = two_column_lookup([(label, "개인식별번호"), (" 김지현 ", "TEST001")])
            result = lookup.lookup("디지털정책과", "김지현")
            self.assertEqual(result["status"], "matched")
            self.assertEqual(result["employee_id"], "TEST001")
            self.assertEqual(result["match_mode"], "name_only")
            self.assertEqual(lookup.lookup("디지털정책과", "김 지현")["status"], "not_found")

    def test_duplicate_names(self):
        lookup = two_column_lookup([("이름", "개인식별번호"), ("김지현", "TEST001"), ("김지현", "TEST002")])
        result = lookup.lookup("디지털정책과", "김지현")
        self.assertEqual(result["status"], "duplicate")
        self.assertIsNone(result["employee_id"])

    def test_name_not_found(self):
        lookup = two_column_lookup([("이름", "개인식별번호"), ("김지현", "TEST001")])
        self.assertEqual(lookup.lookup("디지털정책과", "문형식")["status"], "not_found")

    def test_missing_id(self):
        lookup = two_column_lookup([("이름", "개인식별번호"), ("김지현", None)])
        result = lookup.lookup("디지털정책과", "김지현")
        self.assertEqual(result["status"], "missing_id")
        self.assertIsNone(result["employee_id"])

    def test_duplicate_precedes_missing_id(self):
        lookup = two_column_lookup([("이름", "개인식별번호"), ("김지현", None), ("김지현", "TEST001")])
        self.assertEqual(lookup.lookup("디지털정책과", "김지현")["status"], "duplicate")

    def test_ambiguous_name_headers_rejected(self):
        with self.assertRaises(ValueError):
            two_column_lookup([("이름", "성명", "개인식별번호")])


class EmployeeTests(unittest.TestCase):
    def setUp(self):
        self.lookup = EmployeeLookup(FIXTURES / "employees.xlsx")

    def test_exact_match(self):
        result = self.lookup.lookup("디지털정책과", "문형식")
        self.assertEqual(result["status"], "matched")
        self.assertEqual(result["employee_id"], "TEST002")

    def test_other_department_does_not_interfere(self):
        self.assertEqual(self.lookup.lookup("디지털정책과", "김지현")["employee_id"], "TEST001")
        self.assertEqual(self.lookup.lookup("경제정책과", "김지현")["employee_id"], "TEST003")

    def test_duplicate(self):
        result = EmployeeLookup(FIXTURES / "employees_duplicate.xlsx").lookup("디지털정책과", "김지현")
        self.assertEqual(result["status"], "duplicate")
        self.assertIsNone(result["employee_id"])
        self.assertEqual(len(result["candidates"]), 2)

    def test_not_found(self):
        result = self.lookup.lookup("디지털정책과", "없는테스트직원")
        self.assertEqual(result["status"], "not_found")
        self.assertIsNone(result["employee_id"])
        self.assertEqual(self.lookup.lookup("없는부서", "김지현")["status"], "not_found")

    def test_missing_id(self):
        result = self.lookup.lookup("디지털정책과", "테스트빈번호")
        self.assertEqual(result["status"], "missing_id")
        self.assertIsNone(result["employee_id"])

    def test_no_file(self):
        for path in (None, FIXTURES / "존재하지않음.xlsx"):
            result = EmployeeLookup(path).lookup("디지털정책과", "김지현")
            self.assertEqual(result["status"], "no_file")
            self.assertIsNone(result["employee_id"])

    def test_column_mapping(self):
        columns = {"department": "소속부서", "name": "직원명", "employee_id": "내부번호"}
        lookup = EmployeeLookup(FIXTURES / "employees_mapping.xlsx", columns=columns)
        self.assertEqual(lookup.lookup("디지털정책과", "김지현")["employee_id"], "TEST001")
        with self.assertRaises(ValueError):
            EmployeeLookup(FIXTURES / "employees_mapping.xlsx")

    def test_no_fuzzy_matching(self):
        self.assertEqual(self.lookup.lookup(" 디지털정책과 ", " 김지현 ")["status"], "matched")
        self.assertEqual(self.lookup.lookup("디지털 정책과", "김지현")["status"], "not_found")

    def test_only_matched_can_supply_id(self):
        attendee = {"department": "디지털정책과", "name": "김지현"}
        for status in ("duplicate", "not_found", "no_file", "missing_id"):
            attendee["employee_lookup"] = {"status": status, "employee_id": "TEST999"}
            self.assertIsNone(matched_employee_id(attendee))
        attendee["employee_lookup"] = self.lookup.lookup("경제정책과", "김지현")
        with self.assertRaises(ValueError):
            matched_employee_id(attendee)


class EmployeeGenerationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template = next(ROOT.glob("*작업용.xlsx"))
        cls.notice = parse_notice(next(ROOT.glob("*.pdf")))
        cls.roster = parse_attendees(next(ROOT.glob("*명단.xlsx")))
        cls.reference = load_reference(cls.template, cls.notice["type_raw"])

    def roster_for(self, file):
        roster = deepcopy(self.roster)
        lookup = EmployeeLookup(file)
        for attendee in roster["attendees"]:
            attendee["employee_lookup"] = lookup.lookup(attendee["department"], attendee["name"])
        return roster

    def test_excel_ids_and_preserved_structure(self):
        for file, expected in ((None, [None, None]),
                               (FIXTURES / "employees.xlsx", ["TEST001", "TEST002"]),
                               (FIXTURES / "employees_duplicate.xlsx", [None, "TEST002"])):
            with self.subTest(file=file):
                roster = self.roster_for(file)
                payload = build_excel(self.template, self.notice, roster["attendees"], self.reference["official"], "N", "N", "0")
                verify_excel(self.template, payload, self.notice, roster["attendees"], self.reference["official"], "N", "0")
                workbook = load_workbook(BytesIO(payload), data_only=False)
                try:
                    self.assertEqual([workbook["Sheet1"][f"B{r}"].value for r in (2, 3)], expected)
                    for row in (2, 3):
                        self.assertEqual(workbook["Sheet1"][f"B{row}"].data_type, "s" if expected[row-2] else "n")
                finally:
                    workbook.close()

    def test_all_unmatched_statuses_leave_blank(self):
        for status in ("duplicate", "not_found", "no_file", "missing_id"):
            roster = deepcopy(self.roster)
            for attendee in roster["attendees"]:
                attendee["employee_lookup"] = {"status": status, "employee_id": None}
            payload = build_excel(self.template, self.notice, roster["attendees"], self.reference["official"], "N", "N", "0")
            verify_excel(self.template, payload, self.notice, roster["attendees"], self.reference["official"], "N", "0")

    def test_disk_generation_no_file_and_collision(self):
        roster = self.roster_for(None)
        with TemporaryDirectory() as directory, redirect_stdout(StringIO()):
            for _ in range(2):
                with patch("builtins.input", side_effect=["", "", "0"]):
                    confirm_and_generate(self.template, directory, self.notice, roster, self.reference, 1)
            outputs = list(Path(directory).glob("*.xlsx"))
            self.assertEqual(len(outputs), 2)
            for file in outputs:
                verify_excel(self.template, file.read_bytes(), self.notice, roster["attendees"], self.reference["official"], "N", "0")

    def test_cli_generate_with_virtual_employees(self):
        with TemporaryDirectory() as directory:
            for source in list(ROOT.glob("*.pdf")) + list(ROOT.glob("*.xlsx")):
                shutil.copyfile(source, Path(directory) / source.name)
            run = subprocess.run(
                [sys.executable, "-B", str(ROOT / "src" / "preview.py"),
                 "--input-dir", directory, "--generate", "--employee-file", str(FIXTURES / "employees.xlsx")],
                input="N\nN\n0\n", text=True, capture_output=True, encoding="utf-8",
            )
            self.assertEqual(run.returncode, 0, run.stderr)
            self.assertIn("TEST001", run.stdout)
            self.assertIn("매칭 2명 / 빈셀 0명", run.stdout)
            files = list((Path(directory) / "output").glob("*.xlsx"))
            self.assertEqual(len(files), 1)
            roster = self.roster_for(FIXTURES / "employees.xlsx")
            verify_excel(self.template, files[0].read_bytes(), self.notice, roster["attendees"], self.reference["official"], "N", "0")
