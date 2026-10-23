"""로컬 HTTP API와 실제 샘플을 이용한 웹 업무 흐름 테스트."""
import hashlib
from io import BytesIO
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.app import create_app
from src import settings


class WebTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template = next(ROOT.glob("*작업용.xlsx"))
        cls.pdf = next(ROOT.glob("*.pdf"))
        cls.attendees = next(ROOT.glob("*명단.xlsx"))
        cls.employees = ROOT / "tests/fixtures/employees.xlsx"

    def setUp(self):
        self.directory = TemporaryDirectory()
        self.output = Path(self.directory.name)
        self.client = TestClient(create_app(self.template, self.output))
        self.client.__enter__()
        self.client.headers["X-Requested-With"] = "LearningRecord"
        self.hashes = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.template, self.pdf, self.attendees, self.employees)}

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.directory.cleanup()
        self.assertEqual(self.hashes, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.hashes})

    def analyze(self, employees=True, employee_path=None, department="디지털정책과"):
        files = {"notice": (self.pdf.name, self.pdf.read_bytes(), "application/pdf"),
                 "attendees": (self.attendees.name, self.attendees.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        if employees:
            file = employee_path or self.employees
            files["employees"] = (file.name, file.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return self.client.post("/api/analyze", files=files, data={"department": department})

    def generate(self, analysis, **choices):
        return self.client.post("/api/generate", json={"analysis_id": analysis["analysis_id"], "required": "N", "welfare": "N", "participation": "0", **choices})

    def test_home_and_assets(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("서무_상시학습 등록 자동화", response.text)
        self.assertEqual(response.headers["cache-control"], "no-store")
        for url in ("/static/style.css", "/static/app.js", "/api/config"):
            self.assertEqual(self.client.get(url).status_code, 200)

    def test_upload_analysis_two_attendees_and_ids(self):
        response = self.analyze()
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual([a["name"] for a in data["roster"]["attendees"]], ["김지현", "문형식"])
        self.assertEqual([a["employee_lookup"]["employee_id"] for a in data["roster"]["attendees"]], ["TEST001", "TEST002"])
        self.assertTrue(data["can_generate"])
        self.assertEqual(data["warnings"], [])

    def test_generate_download_choices_formulas_and_blanks(self):
        data = self.analyze().json()
        response = self.generate(data, required="Y", participation="1")
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual((result["count"], result["matched"], result["needs_review"]), (2, 2, 0))
        downloaded = self.client.get(result["download_url"])
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn("attachment", downloaded.headers["content-disposition"])
        self.assertEqual(downloaded.content, (self.output / result["filename"]).read_bytes())
        workbook = load_workbook(BytesIO(downloaded.content), data_only=False)
        try:
            sheet = workbook["Sheet1"]
            self.assertEqual(sheet.max_row, 3)
            for row, name, employee_id in ((2, "김지현", "TEST001"), (3, "문형식", "TEST002")):
                self.assertEqual(sheet[f"A{row}"].value, name)
                self.assertEqual(sheet[f"B{row}"].value, employee_id)
                self.assertEqual(sheet[f"D{row}"].value, f'=INDIRECT("유형코드"&MATCH(C{row},교육유형명,0))')
                self.assertEqual(sheet[f"E{row}"].value, "Y")
                self.assertEqual(sheet[f"K{row}"].value, 1)
                self.assertEqual(sheet[f"N{row}"].value, "N")
                self.assertIsNone(sheet[f"O{row}"].value)
                self.assertIsNone(sheet[f"P{row}"].value)
            self.assertEqual(workbook["교육유형"].sheet_state, "hidden")
            self.assertEqual(workbook["교육기관"].sheet_state, "hidden")
            self.assertTrue(workbook["교육기관"].protection.sheet)
        finally:
            workbook.close()

    def test_no_employee_file_allows_blank_generation(self):
        data = self.analyze(employees=False).json()
        self.assertTrue(data["warnings"])
        self.assertEqual(data["roster"]["attendees"][0]["employee_lookup"]["status"], "no_file")
        result = self.generate(data).json()
        self.assertEqual(result["needs_review"], 2)
        workbook = load_workbook(BytesIO(self.client.get(result["download_url"]).content))
        try:
            self.assertIsNone(workbook["Sheet1"]["B2"].value)
            self.assertIsNone(workbook["Sheet1"]["B3"].value)
        finally:
            workbook.close()

    def test_no_overwrite(self):
        data = self.analyze().json()
        first = self.generate(data).json()
        first_bytes = (self.output / first["filename"]).read_bytes()
        second = self.generate(data).json()
        self.assertNotEqual(first["filename"], second["filename"])
        self.assertEqual((self.output / first["filename"]).read_bytes(), first_bytes)

    def test_missing_required_files_grouped(self):
        response = self.client.post("/api/analyze", data={"department": "디지털정책과"})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(len(response.json()["errors"]), 2)

    def test_wrong_extensions_grouped(self):
        response = self.client.post("/api/analyze", files={"notice": ("공문.txt", b"text"), "attendees": ("명단.xls", b"text")})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(len(response.json()["errors"]), 2)

    def test_corrupt_files_grouped(self):
        response = self.client.post("/api/analyze", files={"notice": ("공문.pdf", b""), "attendees": ("명단.xlsx", b"not an excel")})
        self.assertEqual(response.status_code, 422)

    def test_welfare_yes_blocks_generation(self):
        response = self.generate(self.analyze().json(), welfare="Y")
        self.assertEqual(response.status_code, 422)
        self.assertIn("추가 확인", response.json()["errors"][0])
        self.assertFalse(list(self.output.glob("*.xlsx")))

    def test_invalid_choices(self):
        response = self.generate(self.analyze().json(), participation="9")
        self.assertEqual(response.status_code, 422)

    def test_zero_attendees_blocks_generation(self):
        data = self.analyze(department="없는테스트부서").json()
        self.assertEqual(data["roster"]["attendees"], [])
        self.assertFalse(data["can_generate"])
        self.assertEqual(self.generate(data).status_code, 422)

    def test_other_department_filter(self):
        data = self.analyze(department="경제정책과").json()
        self.assertEqual(len(data["roster"]["attendees"]), 1)
        self.assertEqual(data["roster"]["attendees"][0]["name"], "성지수")

    def test_employee_duplicate_warning_keeps_blank(self):
        data = self.analyze(employee_path=ROOT / "tests/fixtures/employees_duplicate.xlsx").json()
        self.assertIn("동명이인", " ".join(data["warnings"]))
        self.assertEqual(self.generate(data).json()["matched"], 1)

    def test_expired_analysis_and_missing_download(self):
        self.assertEqual(self.client.get("/api/download/unknown").status_code, 404)
        with patch.object(settings, "SESSION_SECONDS", -1):
            data = self.analyze().json()
        self.assertEqual(self.generate(data).status_code, 410)

    def test_cross_origin_requests_rejected(self):
        response = self.client.post("/api/analyze", headers={"origin": "https://example.invalid"})
        self.assertEqual(response.status_code, 403)

    def test_temp_uploads_removed(self):
        import src.app as module
        with TemporaryDirectory() as directory:
            def temporary(**kwargs):
                return TemporaryDirectory(dir=directory, **kwargs)
            with patch.object(module, "TemporaryDirectory", side_effect=temporary):
                self.assertEqual(self.analyze().status_code, 200)
            self.assertEqual(list(Path(directory).iterdir()), [])
